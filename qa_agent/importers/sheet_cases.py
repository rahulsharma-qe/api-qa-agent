"""Existing cases from Excel or CSV, in either of the two layouts teams keep.

API layout, one row per step (a case's later rows may leave Case ID empty):
    Case ID | Title | Criteria | Priority | Tags | Step | Method | Path | Body | Expected Status |
    Expected Fields | Capture
    Body is JSON; Expected Fields and Capture are `name=value; name=value`.

Manual layout, one row per case, steps in words:
    ID | Title | Preconditions | Steps | Expected Result | Priority | Criteria
    -> a case with steps_text: kept and traced, "needs steps" until a person (or the AI layer) writes them.

Each Excel sheet is one feature (its cases go to cases/<sheet name>.yaml); a CSV file is one feature.
"""

import csv
import json
import re
from pathlib import Path

from qa_agent.importers.common import (CaseImportError, ImportedFeature, ImportResult, parse_pairs, priority_from,
                                       split_list)

HEADERS = {
    "id": ("case id", "id", "test case id", "tc id", "test id"),
    "title": ("title", "test case", "test case title", "summary", "scenario", "name"),
    "criteria": ("criteria", "acceptance criteria", "requirement", "requirements", "ac"),
    "priority": ("priority",),
    "tags": ("tags", "labels"),
    "track": ("track",),
    "method": ("method",),
    "path": ("path", "endpoint"),
    "body": ("body", "request body"),
    "status": ("expected status", "status"),
    "fields": ("expected fields",),
    "capture": ("capture", "save"),
    "preconditions": ("preconditions", "pre-conditions", "precondition"),
    "steps": ("steps", "test steps"),
    "expected": ("expected result", "expected results", "expected"),
}
NUMBERED = re.compile(r"(?:^|\s)(\d+)[.)]\s+")


def read_sheet_cases(path: Path, source_name: str) -> ImportResult:
    """source_name: the file as the project sees it (existing/booking_cases.xlsx), kept in each case."""
    result = ImportResult()
    for sheet_name, rows in _tables(path):
        if not rows:
            continue
        columns = _columns(rows[0])
        kind = "excel" if path.suffix.lower() == ".xlsx" else "csv"
        where = f"{path.name}" + (f", sheet '{sheet_name}'" if kind == "excel" else "")
        if "method" in columns and "path" in columns:
            feature = _api_layout(rows, columns, f"{kind}-api", source_name, sheet_name, where, result)
        elif "steps" in columns:
            feature = _manual_layout(rows, columns, f"{kind}-manual", source_name, sheet_name, where, result)
        else:
            raise CaseImportError(f"{where}: the header has neither Method and Path (API layout) nor Steps (manual "
                                  f"layout). Found: {', '.join(str(cell) for cell in rows[0] if cell)}")
        feature.feature = sheet_name if kind == "excel" else path.stem.replace("_", " ")
        result.features.append(feature)
    return result


def _tables(path: Path) -> list[tuple[str, list[list]]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return [(path.stem, [row for row in csv.reader(handle)])]
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]
    raise CaseImportError(f"{path}: use a .xlsx or .csv file")


def _columns(header: list) -> dict[str, int]:
    names = [str(cell or "").strip().lower() for cell in header]
    found = {}
    for key, aliases in HEADERS.items():
        for alias in aliases:
            if alias in names:
                found[key] = names.index(alias)
                break
    return found


def _cell(row: list, columns: dict, key: str) -> str:
    index = columns.get(key)
    if index is None or index >= len(row) or row[index] is None:
        return ""
    value = row[index]
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _base_case(row, columns, case_id, source, row_number, where, result) -> dict:
    priority, note = priority_from(_cell(row, columns, "priority"))
    if note:
        result.notes.append(f"{where}, row {row_number}: {note}")
    case = {"id": case_id, "title": _cell(row, columns, "title") or case_id,
            "criteria": split_list(_cell(row, columns, "criteria")), "priority": priority,
            "track": _cell(row, columns, "track") or "functional"}
    tags = split_list(_cell(row, columns, "tags"))
    if tags:
        case["tags"] = tags
    case["source"] = {**source, "row": row_number}
    return case


def _api_layout(rows, columns, layout, source_name, sheet_name, where, result) -> ImportedFeature:
    feature = ImportedFeature(sheet_name)
    source = {"format": layout, "file": source_name}
    if layout.startswith("excel"):
        source["sheet"] = sheet_name
    current = None
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(cell or "").strip() for cell in row):
            continue
        case_id = _cell(row, columns, "id")
        if case_id:
            current = _base_case(row, columns, case_id, source, row_number, where, result)
            current["steps"] = []
            feature.cases.append(current)
        elif current is None:
            result.notes.append(f"{where}, row {row_number}: no Case ID and no case above it; row skipped")
            continue
        step = _api_step(row, columns, f"{where}, row {row_number}", result)
        if step:
            current["steps"].append(step)
    return feature


def _api_step(row, columns, where, result):
    method, path = _cell(row, columns, "method").upper(), _cell(row, columns, "path")
    if not method or not path:
        result.notes.append(f"{where}: no Method or Path; step skipped")
        return None
    request = {"method": method}
    route, _, query = path.partition("?")
    request["path"] = route
    if query:
        request["query"] = {name: value for name, _, value in (part.partition("=") for part in query.split("&"))}
    body = _cell(row, columns, "body")
    if body:
        try:
            request["body"] = json.loads(body)
        except json.JSONDecodeError:
            result.notes.append(f"{where}: Body is not JSON; kept as text, check it by hand")
            request["body"] = body
    expect = {}
    status = _cell(row, columns, "status")
    if status:
        expect["status"] = int(status) if status.isdigit() else status
    fields = parse_pairs(_cell(row, columns, "fields"))
    if fields:
        expect["fields"] = fields
    step = {"request": request, "expect": expect}
    capture = {name: str(value) for name, value in parse_pairs(_cell(row, columns, "capture")).items()}
    if capture:
        step["capture"] = capture
    return step


def _manual_layout(rows, columns, layout, source_name, sheet_name, where, result) -> ImportedFeature:
    feature = ImportedFeature(sheet_name)
    source = {"format": layout, "file": source_name}
    if layout.startswith("excel"):
        source["sheet"] = sheet_name
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(cell or "").strip() for cell in row):
            continue
        case_id = _cell(row, columns, "id")
        if not case_id:
            result.notes.append(f"{where}, row {row_number}: no ID; row skipped")
            continue
        case = _base_case(row, columns, case_id, source, row_number, where, result)
        lines = []
        if _cell(row, columns, "preconditions"):
            lines.append(f"Before: {_cell(row, columns, 'preconditions')}")
        lines += _step_lines(_cell(row, columns, "steps"))
        if _cell(row, columns, "expected"):
            lines.append(f"Expected: {_cell(row, columns, 'expected')}")
        case["steps_text"] = lines
        feature.cases.append(case)
    return feature


def _step_lines(text: str) -> list[str]:
    """'1. Make a booking 2. Cancel it' or one step per line -> ['Make a booking', 'Cancel it']."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) == 1 and len(NUMBERED.findall(lines[0])) > 1:
        lines = [part.strip() for part in NUMBERED.split(lines[0])[2::2]]
    return [NUMBERED.sub("", line, count=1).strip() if NUMBERED.match(line) else line for line in lines]
