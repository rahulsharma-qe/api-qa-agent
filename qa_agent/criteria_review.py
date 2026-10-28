"""The review sheet's Criteria and Questions tabs: a person approves drafted criteria and answers questions.

`qa-agent review export` adds them when acceptance-criteria-drafts.yaml has something waiting:

    Criteria    each drafted criterion not yet approved or rejected: its text (editable), the words it
                came from (locked), and Decision, Reviewer, Date, Comment, as for cases
    Questions   each open question: its kind, the words it came from, and Answer, Answered by, Date

`qa-agent review apply` applies them: a criterion's decision, only if its text in the file is still what
the sheet was made from; an answer, with who gave it and when.
"""

import datetime

from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from qa_agent.criteria_drafts import CriteriaDrafts
from qa_agent.project import TeamProject
from qa_agent.review_common import (DATE, DECIDE_FILL, DECISIONS, HEADER_FILL, LOCKED_FILL, open_workbook, row_problem,
                                    text_value)

CRITERIA_SHEET, QUESTIONS_SHEET = "Criteria", "Questions"
CRITERIA_COLUMNS = [  # header, width, editable
    ("ID", 9, False), ("Status", 14, False), ("Criterion", 60, True), ("From the requirement", 50, False),
    ("Requirement", 26, False), ("Decision", 16, True), ("Reviewer", 14, True), ("Date", 12, True),
    ("Comment", 40, True), ("Text at export", 10, False)]
QUESTION_COLUMNS = [
    ("ID", 7, False), ("Kind", 14, False), ("Question", 60, False), ("From the requirement", 50, False),
    ("Answer", 50, True), ("Answered by", 16, True), ("Date", 12, True)]
DECIDE = ("Decision", "Reviewer", "Date", "Comment", "Answer", "Answered by")


def pending(project: TeamProject) -> tuple[list, list]:
    """(criteria waiting for a decision, open questions)."""
    drafts = CriteriaDrafts.load(project)
    waiting = [entry for entry in drafts.criteria if entry.get("status", "draft") in ("draft", "changes_needed")]
    return waiting, drafts.open_questions()


def add_criteria_sheets(workbook, project: TeamProject) -> None:
    criteria, questions = pending(project)
    if criteria:
        sheet = workbook.create_sheet(CRITERIA_SHEET, 1)
        sheet.append([name for name, _, _ in CRITERIA_COLUMNS])
        for entry in criteria:
            source = entry.get("source") or {}
            sheet.append([entry["id"], entry.get("status", "draft"), entry.get("text", ""), source.get("quote", ""),
                          source.get("requirement", ""), "", "", "", (entry.get("review") or {}).get("comment", ""),
                          entry.get("text", "")])
        _format(sheet, CRITERIA_COLUMNS, len(criteria))
        sheet.column_dimensions["J"].hidden = True
        choice = DataValidation(type="list", formula1='"Approve,Changes needed,Reject"', allow_blank=True)
        sheet.add_data_validation(choice)
        choice.add(f"F2:F{len(criteria) + 1}")
    if questions:
        sheet = workbook.create_sheet(QUESTIONS_SHEET, 2 if criteria else 1)
        sheet.append([name for name, _, _ in QUESTION_COLUMNS])
        for question in questions:
            sheet.append([question.get("id"), question.get("kind"), question.get("text"), question.get("quote"),
                          "", "", ""])
        _format(sheet, QUESTION_COLUMNS, len(questions))


def apply_criteria_sheets(path, project: TeamProject, result) -> None:
    book = open_workbook(path)
    drafts = CriteriaDrafts.load(project)
    by_id = {str(entry.get("id")): entry for entry in drafts.criteria}
    questions = {str(question.get("id")): question for question in drafts.questions}
    changed = False
    for row in _rows(book, CRITERIA_SHEET):
        if not row.get("Decision"):
            continue
        entry = by_id.get(row["ID"])
        problem = row_problem(row) or ("no such drafted criterion" if entry is None else "") or (
            "" if entry.get("text", "") == row["Text at export"] else
            "the criterion changed after this sheet was made: export the sheet again") or (
            "" if row["Criterion"] else "Criterion is empty")
        if problem:
            result.refused.append(f"{row['ID']}: {problem}")
            continue
        status = DECISIONS[row["Decision"].lower()]
        date = row["Date"] or datetime.date.today().isoformat()
        entry.update(text=row["Criterion"], status=status)
        entry["review"] = {"by": row["Reviewer"], "date": date, "comment": row["Comment"]}
        result.applied.append(f"{row['ID']}: {status.replace('_', ' ')} by {row['Reviewer']}, {date}")
        changed = True
    for row in _rows(book, QUESTIONS_SHEET):
        if not row.get("Answer"):
            continue
        question = questions.get(row["ID"])
        if question is None or question.get("status") == "answered":
            result.refused.append(f"{row['ID']}: no open question with this id")
            continue
        if not row.get("Answered by"):
            result.refused.append(f"{row['ID']}: Answered by is empty: an answer needs a name")
            continue
        if row.get("Date") and not DATE.match(row["Date"]):
            result.refused.append(f"{row['ID']}: Date '{row['Date']}' (write it as 2026-10-24, or leave it empty for today)")
            continue
        answered_on = row.get("Date") or datetime.date.today().isoformat()
        question.update(status="answered", answer=row["Answer"], answered_by=row["Answered by"], answered_on=answered_on)
        result.applied.append(f"{row['ID']}: answered by {row['Answered by']}, {answered_on}")
        changed = True
    if changed:
        drafts.save()


def _rows(book, name: str) -> list[dict]:
    if name not in book.sheetnames:
        return []
    rows = list(book[name].iter_rows(values_only=True))
    header = [str(cell or "") for cell in rows[0]] if rows else []
    return [{column: text_value(value) for column, value in zip(header, values)} for values in rows[1:]]


def _format(sheet, columns, count: int) -> None:
    for index, (name, width, editable) in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
        header = sheet.cell(row=1, column=index)
        header.font, header.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for row in range(2, count + 2):
            cell = sheet.cell(row=row, column=index)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.protection = Protection(locked=not editable)
            if name in DECIDE:
                cell.fill = DECIDE_FILL
            elif not editable:
                cell.fill = LOCKED_FILL
    sheet.freeze_panes = "B2"
    sheet.protection.sheet = True
    sheet.protection.formatColumns = False
    sheet.protection.formatRows = False
