"""`qa-agent trace`: criteria -> cases -> tests -> results, and how much is covered.

    qa-agent trace --results reports/junit.xml target/surefire-reports

Only approved cases count towards coverage; a case in words or a draft is listed, not counted.
Coverage is measured and reported, never decided here: drafting cases for the gaps is the AI
layer's job. `--check` fails only when an approved case has no test.

    Criteria        each criterion has at least one approved case
    Operations      each operation in api/api.yaml has one (not_tested: <reason> leaves it out)
    Response codes  each documented code of each operation is expected by an approved step;
                    401 counts once per service, 422 only where the operation takes a body
    Cases -> tests  each approved case has its test in every framework
    Case types      approved cases by type (positive, negative, auth, edge: see case_types.py), and the
                    criteria whose approved cases are all positive: the happy path only
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from qa_agent.api_model import ApiModel, ApiModelError, Operation, load_api_model
from qa_agent.checks import plural
from qa_agent.case_types import NOT_HAPPY, TYPES, case_types, expected_statuses, label
from qa_agent.cases import Case, load_case_file
from qa_agent.project import TeamProject
from qa_agent.renderers import RENDERERS
from qa_agent.results import read_results
from qa_agent.review_sheet import natural_order

TAKES_A_BODY = ("POST", "PUT", "PATCH")
FLOORED = ("Criteria", "Operations")
TEST_NAME = {"pytest": re.compile(r"^def (test_\w+)\(", re.MULTILINE),
             "restassured": re.compile(r"^\s+public void (\w+)\(", re.MULTILINE)}


@dataclass
class Coverage:
    name: str
    covered: int
    total: int
    missing: list = field(default_factory=list)
    note: str = ""

    @property
    def percent(self) -> float:
        return 100.0 * self.covered / self.total if self.total else 100.0

    def as_line(self) -> str:
        if self.note:
            return f"{self.name:<15} {self.note}"
        line = f"{self.name:<15} {self.covered} of {self.total}"
        return line + (f"   missing: {', '.join(self.missing)}" if self.missing else "")


@dataclass
class TraceRow:
    criterion: str
    case_id: str
    title: str
    review: str
    suite: str
    tests: list
    result: str
    types: str = "-"


@dataclass
class TraceReport:
    team: str
    rows: list = field(default_factory=list)
    coverage: list = field(default_factory=list)
    not_tested: list = field(default_factory=list)        # "POST /test/reset: a test-only route"
    missing_tests: list = field(default_factory=list)     # "TC001 (restassured)"
    below_floor: list = field(default_factory=list)       # "Operations 62% (floor 80%)", with --min-coverage
    orphan_tests: list = field(default_factory=list)
    results_read: dict = field(default_factory=dict)      # framework -> {"passed": n, ...}
    results_left_out: int = 0                             # results of tests no longer generated (older files)
    case_types: dict = field(default_factory=dict)        # "positive" -> how many approved cases
    approved: int = 0
    untyped: list = field(default_factory=list)           # approved cases with no type yet: no steps, status TBD
    happy_path_only: list = field(default_factory=list)   # criteria whose approved cases are all positive

    def types_line(self) -> str:
        counts = " · ".join(f"{self.case_types.get(kind, 0)} {kind}" for kind in TYPES)
        line = f"{self.approved} approved: {counts} (a case can be of more than one type)"
        return line + (f"   no type yet: {', '.join(self.untyped)}" if self.untyped else "")

    def results_line(self) -> str:
        read = " · ".join(f"{framework} " + ", ".join(f"{n} {kind}" for kind, n in counts.items())
                          for framework, counts in self.results_read.items())
        left_out = (f" · {plural(self.results_left_out, 'result')} of tests that no longer exist left out"
                    if self.results_left_out else "")
        return (read or "none given (--results reports/junit.xml target/surefire-reports)") + left_out

    def as_text(self) -> str:
        lines = [f"qa-agent trace: {self.team}", ""] + [item.as_line() for item in self.coverage]
        lines.append(f"{'':<15} ({COVERED_MEANS})")
        if self.not_tested:
            lines.append(f"{'Not tested':<15} on purpose: " + "; ".join(self.not_tested))
        if self.approved:
            lines.append(f"{'Case types':<15} {self.types_line()}")
        if self.happy_path_only:
            lines.append(f"{'Happy path only':<15} {', '.join(self.happy_path_only)}: no negative, auth or edge case")
        lines.append(f"{'Results':<15} {self.results_line()}")
        if self.missing_tests:
            lines.append(f"{'No test':<15} " + ", ".join(self.missing_tests) + ": run qa-agent generate")
        if self.orphan_tests:
            lines.append(f"{'No case':<15} " + ", ".join(self.orphan_tests))
        if self.below_floor:
            lines.append(f"{'Below floor':<15} " + ", ".join(self.below_floor))
        return "\n".join(lines)

    def apply_floor(self, minimum: float) -> None:
        """--min-coverage: criteria and operations coverage must reach it (response codes are reported only)."""
        for item in self.coverage:
            if item.name in FLOORED and not item.note and item.percent < minimum:
                self.below_floor.append(f"{item.name} {item.percent:.0f}% (floor {minimum:.0f}%)")

    def as_markdown(self) -> str:
        lines = [f"# Traceability matrix: {self.team}", "", "| Coverage | Covered | Missing |", "|---|---|---|"]
        for item in self.coverage:
            lines.append(f"| {item.name} | {item.note or f'{item.covered} of {item.total}'} | "
                         f"{', '.join(item.missing) or '-'} |")
        lines += ["", COVERED_MEANS[0].upper() + COVERED_MEANS[1:] + ".",
                  "", f"Results: {self.results_line()}"]
        if self.not_tested:
            lines += ["", "Not tested, on purpose: " + "; ".join(self.not_tested)]
        if self.approved:
            lines += ["", f"Case types: {self.types_line()}"]
        if self.happy_path_only:
            lines += ["", f"Happy path only: {', '.join(self.happy_path_only)} (no negative, auth or edge case)"]
        lines += ["", "| Criterion | Case | Title | Type | Review | Suite | Tests | Latest result |",
                  "|---|---|---|---|---|---|---|---|"]
        for row in self.rows:
            lines.append(f"| {row.criterion} | {row.case_id or '-'} | {row.title} | {row.types} | {row.review} | "
                         f"{row.suite} | {'<br>'.join(f'`{test}`' for test in row.tests) or '-'} | {row.result} |")
        return "\n".join(lines) + "\n"

    def write_excel(self, path: Path) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        matrix = book.active
        matrix.title = "Traceability matrix"
        matrix.append(["Criterion", "Case", "Title", "Type", "Review", "Suite", "Tests", "Latest result"])
        for row in self.rows:
            matrix.append([row.criterion, row.case_id, row.title, row.types, row.review, row.suite,
                           "\n".join(row.tests), row.result])
        coverage = book.create_sheet("Coverage")
        coverage.append(["Coverage", "Covered", "Of", "Missing"])
        for item in self.coverage:
            coverage.append([item.name, item.covered, item.total, item.note or "\n".join(item.missing)])
        coverage.append(["", COVERED_MEANS[0].upper() + COVERED_MEANS[1:]])
        coverage.append(["Results", self.results_line()])
        if self.approved:
            coverage.append(["Case types", self.types_line()])
        if self.happy_path_only:
            coverage.append(["Happy path only", ", ".join(self.happy_path_only)])
        for sheet, widths in ((matrix, (11, 10, 44, 16, 18, 12, 60, 14)), (coverage, (16, 9, 6, 70))):
            for index, width in enumerate(widths):
                sheet.column_dimensions[chr(65 + index)].width = width
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for line in sheet.iter_rows(min_row=2):
                for cell in line:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)


COVERED_MEANS = "covered means an approved case exists; whether its tests pass is under Results"


def build_trace(project: TeamProject, result_paths: Optional[list] = None) -> TraceReport:
    report = TraceReport(project.team)
    criteria = project.read_criteria()
    cases = sorted((case for path in project.case_files() for case in load_case_file(path).cases),
                   key=lambda case: natural_order(case.id))
    approved = [case for case in cases if case.is_approved]
    written = {framework: _written_tests(project, framework) for framework in project.frameworks}
    results = read_results(result_paths or [])
    outcomes = results.outcomes
    expected_names = set()
    tests_by_case = {}
    for case in approved:
        tests_by_case[case.id] = []
        for framework in project.frameworks:
            for name in RENDERERS[framework].test_names(case):
                expected_names.add(name)
                if name in written[framework]:
                    tests_by_case[case.id].append(f"{framework}: {written[framework][name]}")
                else:
                    report.missing_tests.append(f"{case.id} ({framework})")
    report.orphan_tests = sorted(f"{framework}: {label}" for framework in project.frameworks
                                 for name, label in written[framework].items() if name not in expected_names)
    report.results_read, report.results_left_out = _current_results(results, expected_names)
    types = {case.id: case_types(case, project) for case in approved}
    report.rows = _rows(cases, criteria, tests_by_case, outcomes, project, types)
    _count_types(report, approved, criteria, types)
    report.coverage = [_criteria_coverage(criteria, cases)]
    api = read_api(project)
    if api is None:
        report.coverage.append(Coverage("Operations", 0, 0, note="no api/api.yaml: qa-agent import api <file>"))
    else:
        report.coverage += _api_coverage(api, approved, project, report)
    with_tests = sum(1 for case in approved if not any(item.startswith(case.id + " ") for item in report.missing_tests))
    report.coverage.append(Coverage("Cases to tests", with_tests, len(approved),
                                    [item for item in report.missing_tests]))
    return report


def _current_results(results, expected_names: set) -> tuple:
    """Counts per framework for the tests that exist now; results of others (an older file) are left out."""
    counts, left_out = {}, 0
    for test in results.runs:
        if test.name not in expected_names:
            left_out += 1
            continue
        framework = "pytest" if test.name.startswith("test_") else "restassured"
        counts.setdefault(framework, {})
        counts[framework][test.outcome] = counts[framework].get(test.outcome, 0) + 1
    return counts, left_out


def _written_tests(project: TeamProject, framework: str) -> dict:
    """Test name -> 'file::name', for every test in the framework's generated folder."""
    renderer = RENDERERS[framework]
    folder = project.root / renderer.output_dir
    found = {}
    for path in sorted(folder.glob("*.py" if framework == "pytest" else "*.java")) if folder.is_dir() else []:
        text = path.read_text(encoding="utf-8")
        if not text.startswith(renderer.header_mark):
            continue
        for name in TEST_NAME[framework].findall(text):        # data providers return Object[][], not void
            found[name] = f"{path.name}::{name}"
    return found


def _result(case: Case, project: TeamProject, outcomes: dict) -> str:
    """Per framework: 'pytest passed 3/3 · restassured not run'. A data-driven case has a run per row."""
    labels = []
    for framework in project.frameworks:
        seen = [outcome for name in RENDERERS[framework].test_names(case) for outcome in outcomes.get(name, [])]
        if not seen:
            label = "not run"
        elif "failed" in seen:
            label = f"failed {seen.count('failed')}/{len(seen)}"
        elif all(outcome == "skipped" for outcome in seen):
            label = "skipped"
        else:
            label = f"passed {seen.count('passed')}/{len(seen)}"
        labels.append(f"{framework} {label}" if len(project.frameworks) > 1 else label)
    return " · ".join(labels)


def _review_label(case: Case) -> str:
    if case.changed_after_approval:
        return "changed after approval"
    if case.needs_steps:
        return f"{case.review.status.replace('_', ' ')} (needs steps)"
    return case.review.status.replace("_", " ")


def _count_types(report: TraceReport, approved: list, criteria: dict, types: dict) -> None:
    report.approved = len(approved)
    report.case_types = {kind: sum(1 for case in approved if kind in types[case.id]) for kind in TYPES}
    report.untyped = [case.id for case in approved if not types[case.id]]
    for criterion in criteria:
        seen = set().union(*(types[case.id] for case in approved if criterion in case.criteria))
        if seen and not seen & set(NOT_HAPPY):
            report.happy_path_only.append(criterion)


def _rows(cases: list, criteria: dict, tests_by_case: dict, outcomes: dict, project: TeamProject,
          types: dict) -> list:
    found = []
    for criterion in list(criteria) + ["(no criterion)"]:
        linked = [case for case in cases if criterion in case.criteria
                  or (criterion == "(no criterion)" and not set(case.criteria) & set(criteria))]
        if not linked and criterion != "(no criterion)":
            found.append(TraceRow(criterion, "", "no case", "-", "-", [], "-"))
        for case in linked:
            tests = tests_by_case.get(case.id, [])
            result = _result(case, project, outcomes) if case.is_approved else "-"
            found.append(TraceRow(criterion, case.id, case.title, _review_label(case), case.suite, tests, result,
                                  label(types.get(case.id, set())) if case.is_approved else "-"))
    return found


def _criteria_coverage(criteria: dict, cases: list) -> Coverage:
    if not criteria:
        return Coverage("Criteria", 0, 0, note="no acceptance-criteria.yaml")
    missing = []
    for criterion in criteria:
        linked = [case for case in cases if criterion in case.criteria and case.review.status != "rejected"]
        if any(case.is_approved for case in linked):
            continue
        if not linked:
            missing.append(f"{criterion} (no case)")
        else:
            missing.append(f"{criterion} (not approved: {', '.join(case.id for case in linked)})")
    return Coverage("Criteria", len(criteria) - len(missing), len(criteria), missing)


def read_api(project: TeamProject) -> Optional[ApiModel]:
    if not project.api_model_path.is_file():
        return None
    try:
        return load_api_model(project.api_model_path)
    except ApiModelError:
        return None


@dataclass
class ApiItem:
    """One thing the API coverage counts: an operation, or one of its response codes."""
    label: str
    covered: bool
    service: str = ""
    operation: Optional[Operation] = None
    code: Optional[int] = None


def api_items(api: ApiModel, approved: list, project: TeamProject) -> tuple[list, list, list]:
    """(operations, response codes, not tested): what the coverage counts, and what approved cases call.
    401 counts once per service (its item has no operation); 422 only where the operation takes a body."""
    called: dict = {}          # (service, method, path) -> statuses expected by approved steps
    for case in approved:
        for step in case.steps:
            service = step.request.service or project.main_service
            operation = api.find(service, step.request.method, step.request.path)
            if operation is not None:
                called.setdefault((service, operation.method, operation.path), set()).update(
                    expected_statuses(step.expect.get("status"), case, project))
    operations, codes, not_tested, services_with_401 = [], [], [], {}
    for service, service_operations in sorted(api.services.items()):
        for operation in service_operations:
            label = f"{operation.method} {operation.path}" + ("" if service == project.main_service else f" ({service})")
            if operation.not_tested:
                not_tested.append(f"{label}: {operation.not_tested}")
                continue
            statuses = called.get((service, operation.method, operation.path))
            operations.append(ApiItem(label, statuses is not None, service, operation))
            for code in operation.responses:
                if code == 401:
                    services_with_401.setdefault(service, False)
                    services_with_401[service] |= 401 in (statuses or set())
                elif code != 422 or operation.method in TAKES_A_BODY:
                    codes.append(ApiItem(f"{label} {code}", code in (statuses or set()), service, operation, code))
    for service, seen in services_with_401.items():
        codes.append(ApiItem(f"401 ({service}, any operation)", seen, service, None, 401))
    return operations, codes, not_tested


def _api_coverage(api: ApiModel, approved: list, project: TeamProject, report: TraceReport) -> list:
    operations, codes, not_tested = api_items(api, approved, project)
    report.not_tested += not_tested
    return [_coverage("Operations", [(item.label, item.covered) for item in operations]),
            _coverage("Response codes", [(item.label, item.covered) for item in codes])]


def _coverage(name: str, items: list) -> Coverage:
    return Coverage(name, sum(1 for _, seen in items if seen), len(items), [label for label, seen in items if not seen])
