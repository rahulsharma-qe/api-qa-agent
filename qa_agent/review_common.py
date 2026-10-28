"""What both tabs of the review sheet share: the decisions, the colours, how a decided row is read."""

import datetime
import re

from openpyxl.styles import PatternFill

DECISIONS = {"approve": "approved", "approved": "approved", "changes needed": "changes_needed",
             "reject": "rejected", "rejected": "rejected"}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
LOCKED_FILL = PatternFill("solid", fgColor="F3F4F6")
DECIDE_FILL = PatternFill("solid", fgColor="FEF9C3")
DATE_FORMATS = ("%d-%m-%Y", "%Y-%m-%d")             # 28-10-2026, as people write it; 2026-10-28, as Excel gives it


def row_problem(row: dict) -> str:
    """Why a decided row can't be applied (a case or a criterion), or ''."""
    if row["Decision"].lower() not in DECISIONS:
        return f"Decision '{row['Decision']}' (use Approve, Changes needed or Reject)"
    if not row.get("Reviewer"):
        return "Reviewer is empty: a decision needs a name"
    if DECISIONS[row["Decision"].lower()] != "approved" and not row.get("Comment"):
        return f"{row['Decision']} needs a Comment, so the author knows why"
    if row.get("Date") and not review_date(row["Date"]):
        return f"Date '{row['Date']}' isn't a date (write it as 28-10-2026, or leave it empty for today)"
    return ""


def shown_date(iso: str) -> str:
    """2026-10-28 -> 28-10-2026: every date a person reads is DD-MM-YYYY."""
    try:
        return datetime.date.fromisoformat(iso).strftime("%d-%m-%Y")
    except (TypeError, ValueError):
        return iso


def review_date(text: str) -> str:
    """'28-10-2026' or '2026-10-28' -> '2026-10-28' (kept sortable in the case files); '' when it isn't a real date.
    Empty means today."""
    if not text:
        return datetime.date.today().isoformat()
    for pattern in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text.strip(), pattern).date().isoformat()
        except ValueError:
            continue
    return ""


def text_value(value) -> str:
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "" if value is None else str(value).strip()


def open_workbook(path):
    """An .xlsx a person filled, or a SetupError that says what to do (no zipfile traceback)."""
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    from qa_agent.project import SetupError
    try:
        return load_workbook(path, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise SetupError(f"{path} can't be opened as an Excel sheet ({type(error).__name__}): it may be damaged, or "
                         "saved in another format; save it again as .xlsx (Excel Workbook)") from None
