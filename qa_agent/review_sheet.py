"""The review sheet: a person approves, sends back or rejects each case, in Excel.

`qa-agent review export` writes review/cases.xlsx: one row per case still to review. The steps and
expected results are readable text, locked: steps are changed in the case file, so a sheet can never
produce a case that doesn't parse. The reviewer may change Title, Criteria, Priority, Suite and
Tags, and fills Decision, Reviewer, Date and Comment.

`qa-agent review apply` applies each decided row, only if the case is unchanged since the export
(a hidden column holds its fingerprint). Approve records who and when, with the fingerprint; a case
with check errors, or with steps only in words, can't be approved.
"""

import datetime
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from qa_agent.cases import PRIORITIES, SUITES, Case, load_case_file, save_case_file
from qa_agent.checks import ERROR, check_project, plural
from qa_agent.criteria_review import add_criteria_sheets, apply_criteria_sheets
from qa_agent.review_common import (DECIDE_FILL, DECISIONS, HEADER_FILL, LOCKED_FILL, row_problem,
                                    text_value)
from qa_agent.importers.common import split_list
from qa_agent.project import SetupError, TeamProject

SHEET = "Cases"
COLUMNS = [  # header, width, editable
    ("ID", 12, False), ("Case file", 22, False), ("Status", 14, False), ("Title", 36, True),
    ("Criteria", 12, True), ("Priority", 9, True), ("Suite", 12, True), ("Tags", 12, True),
    ("Steps", 60, False), ("Expected", 44, False), ("Check findings", 44, False),
    ("Decision", 16, True), ("Reviewer", 14, True), ("Date", 12, True), ("Comment", 40, True),
    ("Fingerprint", 14, False)]
HOW_TO = [
    "How to review",
    "",
    "1. Read each case: its steps and expected results (grey, locked), and what qa-agent check found.",
    "2. You may change Title, Criteria, Priority, Suite and Tags. Steps are changed in the case file.",
    "3. Pick a Decision: Approve, Changes needed or Reject. Fill Reviewer; Date is today if left empty.",
    "4. Changes needed and Reject need a Comment, so the author knows why.",
    "5. Save, then run: qa-agent review apply",
    "",
    "Criteria tab (when there is one): drafted acceptance criteria. Approve, send back or reject each, as for",
    "cases; you may reword the criterion. Questions tab: what the requirement leaves open. Write the Answer",
    "and who gave it (Answered by), and its Date (today if left empty); the product owner is often the one to ask.",
    "",
    "Only approved cases become tests. A case edited after this sheet was made is not applied:",
    "export the sheet again. A case with check errors, or with steps only in words, can't be approved.",
]


@dataclass
class ReviewImportResult:
    applied: list = field(default_factory=list)
    refused: list = field(default_factory=list)

    def as_text(self) -> str:
        lines = [f"applied  {line}" for line in self.applied] + [f"refused  {line}" for line in self.refused]
        lines.append(f"{plural(len(self.applied), 'decision')} applied, {len(self.refused)} refused")
        return "\n".join(lines)


def export_review_sheet(project: TeamProject, path: Path, every_case: bool = False) -> int:
    """Writes the sheet; returns the number of cases in it."""
    report = check_project(project)
    rows = [(case_file, case) for case_file in report.case_files for case in case_file.cases
            if every_case or (not case.is_approved and case.review.status != "rejected")]
    rows.sort(key=lambda row: natural_order(row[1].id))                    # TC001, TC002, ... TC010
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append([name for name, _, _ in COLUMNS])
    for case_file, case in rows:
        findings = [f"{f.level}: {f.message[len(case.id):].lstrip(': ')}" for f in report.for_case(case.id)]
        if case.notes:
            findings.append(f"notes: {case.notes}")
        sheet.append([case.id, project.relative(case_file.path), case.review.status, case.title,
                      ", ".join(case.criteria), case.priority, case.suite, ", ".join(case.tags),
                      readable_steps(case), readable_expectations(case), "\n".join(findings),
                      "", "", "", case.review.comment, case.fingerprint()])
    _format(sheet, len(rows))
    add_criteria_sheets(workbook, project)
    guide = workbook.create_sheet("How to review")
    for line in HOW_TO:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return len(rows)


def import_review_sheet(project: TeamProject, path: Path) -> ReviewImportResult:
    result = ReviewImportResult()
    decided = _decided_rows(path, result)
    files = {case_file.path: case_file for case_file in map(load_case_file, project.case_files())}
    cases = {case.id: (case_file, case) for case_file in files.values() for case in case_file.cases}
    accepted = []
    for row in decided:
        found = cases.get(row["ID"])
        if found is None:
            result.refused.append(f"{row['ID']}: no such case in the project")
            continue
        case_file, case = found
        problem = _edit(case, row) if case.fingerprint() == row["Fingerprint"] else \
            "the case changed after this sheet was made: export the sheet again"
        if problem:
            result.refused.append(f"{case.id}: {problem}")
            continue
        accepted.append((case_file, case, row))
    for case_file in {case_file.path: case_file for case_file, _, _ in accepted}.values():
        save_case_file(case_file)
    report = check_project(project)
    for case_file, case, row in accepted:
        status = DECISIONS[row["Decision"].lower()]
        errors = [f for f in report.for_case(case.id) if f.level == ERROR]
        if status == "approved" and (case.needs_steps or errors):
            reason = "it has steps only in words" if case.needs_steps else f"check finds {plural(len(errors), 'error')}"
            result.refused.append(f"{case.id}: can't be approved: {reason}; fix it, then review it again")
            continue
        case.review.status, case.review.by, case.review.date = status, row["Reviewer"], row["Date"]
        case.review.comment = row["Comment"]
        case.review.fingerprint = case.fingerprint() if status == "approved" else ""
        result.applied.append(f"{case.id}: {status.replace('_', ' ')} by {row['Reviewer']}, {row['Date']}")
    for case_file in {case_file.path: case_file for case_file, _, _ in accepted}.values():
        save_case_file(case_file)
    apply_criteria_sheets(path, project, result)
    return result


def readable_steps(case: Case) -> str:
    if case.needs_steps:
        return "\n".join(line for line in case.manual_steps if not _is_expected(line))
    lines = []
    for number, step in enumerate(case.steps, start=1):
        request = step.request
        text = f"{number}. {request.method} {request.path}"
        if request.service:
            text += f" ({request.service})"
        if request.query:
            text += "?" + "&".join(f"{k}={v}" for k, v in request.query.items())
        if request.body is not None:
            text += " " + _short(json.dumps(request.body, ensure_ascii=False))
        if request.auth == "none":
            text += " (no token)"
        if step.capture:
            text += " → remember " + ", ".join(step.capture)
        lines.append(text)
    return "\n".join(lines)


def readable_expectations(case: Case) -> str:
    if case.needs_steps:
        expected = [re.sub(r"^Expected:\s*", "", line) for line in case.manual_steps if _is_expected(line)]
        return "\n".join(expected + ["(in words: needs steps)"]) if expected else \
            "(in words: needs steps; no expected result written)"
    lines = []
    for number, step in enumerate(case.steps, start=1):
        expect = step.expect
        parts = [str(expect["status"])] if "status" in expect else []
        parts += [f"{path} = {json.dumps(value, ensure_ascii=False)}" for path, value in expect.get("fields", {}).items()]
        parts += [f"has {path}" for path in expect.get("exists", [])]
        parts += [f"{path} is {kind}" for path, kind in expect.get("types", {}).items()]
        parts += [f"{path} contains '{text}'" for path, text in expect.get("contains", {}).items()]
        parts += [f"{path} has {count} items" for path, count in expect.get("list_length", {}).items()]
        parts += [f"{path} has an item {_short(json.dumps(m))}" for path, m in expect.get("list_contains", {}).items()]
        parts += [f"{path} has no item {_short(json.dumps(m))}" for path, m in expect.get("list_excludes", {}).items()]
        parts += [f"header {name} = {value}" for name, value in expect.get("headers", {}).items()]
        if "max_response_ms" in expect:
            parts.append(f"within {expect['max_response_ms']} ms")
        lines.append(f"{number}. " + ("; ".join(parts) or "(nothing checked)"))
    return "\n".join(lines)


def _is_expected(line: str) -> bool:
    """'Expected: ...' from a manual sheet, or a Gherkin 'Then ...'."""
    return line.startswith(("Expected:", "Then "))


def natural_order(case_id: str) -> list:
    """TC2 before TC10."""
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", case_id)]


def _short(text: str, limit: int = 120) -> str:
    return text if len(text) <= limit else text[:limit - 1] + "…"


def _format(sheet, count: int) -> None:
    for index, (name, width, editable) in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = width
        header = sheet.cell(row=1, column=index)
        header.font, header.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for row in range(2, count + 2):
            cell = sheet.cell(row=row, column=index)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.protection = Protection(locked=not editable)
            if not editable:
                cell.fill = LOCKED_FILL
            elif name in ("Decision", "Reviewer", "Date", "Comment"):
                cell.fill = DECIDE_FILL
    sheet.column_dimensions["P"].hidden = True
    sheet.freeze_panes = "B2"
    last = count + 1
    for column, values in (("L", "Approve,Changes needed,Reject"), ("F", ",".join(PRIORITIES)),
                           ("G", ",".join(SUITES))):
        choice = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        sheet.add_data_validation(choice)
        choice.add(f"{column}2:{column}{max(last, 2)}")
    sheet.protection.sheet = True
    sheet.protection.formatColumns = False         # the reviewer may still widen columns and rows
    sheet.protection.formatRows = False


def _decided_rows(path: Path, result: ReviewImportResult) -> list[dict]:
    book = load_workbook(path, data_only=True)
    if SHEET not in book.sheetnames:
        raise SetupError(f"{path}: no '{SHEET}' sheet; export the review sheet again")
    rows = list(book[SHEET].iter_rows(values_only=True))
    header = [str(cell or "") for cell in rows[0]] if rows else []
    missing = [name for name, _, _ in COLUMNS if name not in header]
    if missing:
        raise SetupError(f"{path}: columns missing: {', '.join(missing)}; export the review sheet again")
    decided = []
    for number, values in enumerate(rows[1:], start=2):
        row = {name: text_value(value) for name, value in zip(header, values)}
        if not row.get("Decision"):
            continue
        problem = row_problem(row)
        if problem:
            result.refused.append(f"{row.get('ID') or f'row {number}'}: {problem}")
            continue
        row["Date"] = row["Date"] or datetime.date.today().isoformat()
        decided.append(row)
    return decided


def _edit(case: Case, row: dict) -> str:
    """The reviewer's edits to the title, criteria, priority, suite and tags. Returns a problem, if any."""
    if row["Priority"] not in PRIORITIES:
        return f"Priority '{row['Priority']}' (use {', '.join(PRIORITIES)})"
    if row["Suite"] not in SUITES:
        return f"Suite '{row['Suite']}' (use {' or '.join(SUITES)})"
    if not row["Title"]:
        return "Title is empty"
    case.title, case.priority, case.suite = row["Title"], row["Priority"], row["Suite"]
    case.criteria, case.tags = split_list(row["Criteria"]), split_list(row["Tags"])
    return ""
