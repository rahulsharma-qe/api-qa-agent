"""Test data from files, and fresh values so tests never clash.

    load_rows("data/cancellation_boundaries.csv")   -> [{"case": "...", "refund_allowed": True, ...}, ...]
    new_passenger()                                  -> {"passenger_name": "...", "email": "...@testdata.example"}

CSV and Excel cells become numbers, true/false and empty (None) where they look like them; JSON is
read as written (a list of objects, or {"rows": [...]}).
"""

import csv
import json
import secrets
from pathlib import Path
from typing import Any

TEST_EMAIL_DOMAIN = "testdata.example"


def load_rows(path) -> list[dict]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"test data file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        rows = data.get("rows") if isinstance(data, dict) else data
        if not isinstance(rows, list):
            raise ValueError(f"{path}: expected a list of objects, or {{\"rows\": [...]}}")
        return rows
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8") as handle:
            return [{key: _typed(value) for key, value in row.items()} for row in csv.DictReader(handle)]
    if suffix in (".xlsx", ".xlsm"):
        return _excel_rows(path)
    raise ValueError(f"{path}: use .json, .csv or .xlsx")


def new_passenger() -> dict:
    """A passenger no other test is using: a unique name and a test-only email address."""
    tag = secrets.token_hex(3)
    return {"passenger_name": f"Test Passenger {tag}", "email": f"passenger-{tag}@{TEST_EMAIL_DOMAIN}"}


def _excel_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(cell).strip() for cell in rows[0]]
    return [{key: _typed(value) for key, value in zip(header, row)} for row in rows[1:]
            if any(cell is not None for cell in row)]


def _typed(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text == "":
        return None
    if text.lower() in ("true", "false"):
        return text.lower() == "true"
    for convert in (int, float):
        try:
            return convert(text)
        except ValueError:
            pass
    return text
