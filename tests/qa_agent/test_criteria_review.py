"""The review sheet's Criteria and Questions tabs: drafted criteria approved by a person, questions answered."""

import datetime

from openpyxl import load_workbook

from qa_agent.cli import main
from qa_agent.criteria_drafts import CriteriaDrafts
from qa_agent.project import TeamProject
from qa_agent.review_sheet import export_review_sheet, import_review_sheet

DRAFTS = """
criteria:
  - {id: AC-2, text: A booking cancelled within 24 hours gets a full refund., status: draft,
     source: {kind: ai-draft, requirement: requirements/cancel.md, quote: within 24 hours}}
  - {id: AC-3, text: A booking cancelled later gets no refund., status: draft}
  - {id: AC-4, text: Already approved., status: approved}
questions:
  - {id: Q-1, kind: ambiguity, text: "Exactly 24 hours?", quote: within 24 hours, status: open}
  - {id: Q-2, kind: dependency, text: "Payments down?", quote: payments service, status: open}
"""


def sheet_with(make_project, tmp_path):
    root = make_project({"acceptance-criteria.yaml": "criteria:\n  - {id: AC-1, text: A booking can be made.}\n",
                         "acceptance-criteria-drafts.yaml": DRAFTS})
    project = TeamProject.load(root)
    sheet = root / "review" / "cases.xlsx"
    export_review_sheet(project, sheet)
    return project, sheet


def fill(sheet, tab, rows):
    book = load_workbook(sheet)
    header = [cell.value for cell in book[tab][1]]
    for number, values in rows.items():
        for column, value in values.items():
            book[tab].cell(row=number, column=header.index(column) + 1, value=value)
    book.save(sheet)


def test_the_tabs_list_what_waits_for_a_person(make_project, tmp_path):
    _, sheet = sheet_with(make_project, tmp_path)
    book = load_workbook(sheet)
    assert book.sheetnames == ["Cases", "Criteria", "Questions", "How to review"]
    criteria = list(book["Criteria"].iter_rows(values_only=True))
    assert [row[0] for row in criteria[1:]] == ["AC-2", "AC-3"]                     # AC-4 is approved already
    assert criteria[1][3] == "within 24 hours" and book["Criteria"].column_dimensions["J"].hidden
    assert [row[0] for row in list(book["Questions"].iter_rows(values_only=True))[1:]] == ["Q-1", "Q-2"]


def test_decisions_and_answers_are_applied_with_who_and_when(make_project, tmp_path, capsys):
    project, sheet = sheet_with(make_project, tmp_path)
    fill(sheet, "Criteria", {2: {"Decision": "Approve", "Reviewer": "qa-lead", "Date": "2026-10-24",
                                 "Criterion": "A booking cancelled within 24 hours, 24 included, gets a full refund."},
                             3: {"Decision": "Reject", "Reviewer": "qa-lead", "Comment": "covered by AC-1"}})
    fill(sheet, "Questions", {2: {"Answer": "Yes, 24 hours exactly is refunded", "Answered by": "product-owner"}})
    assert main(["--project", str(project.root), "review", "apply"]) == 0
    out = capsys.readouterr().out
    assert "AC-2: approved by qa-lead, 2026-10-24" in out and "Q-1: answered by product-owner" in out
    assert project.read_criteria()["AC-2"].text == "A booking cancelled within 24 hours, 24 included, gets a full refund."
    assert project.read_all_criteria()["AC-3"].status == "rejected"
    drafts = CriteriaDrafts.load(project)
    assert drafts.criteria[0]["review"] == {"by": "qa-lead", "date": "2026-10-24", "comment": ""}
    assert [q["id"] for q in drafts.open_questions()] == ["Q-2"]
    assert drafts.questions[0]["answer"] == "Yes, 24 hours exactly is refunded"


def test_a_row_that_breaks_a_rule_is_refused_and_changes_nothing(make_project, tmp_path):
    project, sheet = sheet_with(make_project, tmp_path)
    fill(sheet, "Criteria", {2: {"Decision": "Approve", "Reviewer": "qa-lead"},
                             3: {"Decision": "Changes needed", "Reviewer": "qa-lead"}})
    fill(sheet, "Questions", {3: {"Answer": "It retries"}})
    drafts_path = project.root / "acceptance-criteria-drafts.yaml"
    drafts_path.write_text(drafts_path.read_text().replace("gets a full refund.", "gets all the money back."))
    refused = "\n".join(import_review_sheet(project, sheet).refused)
    assert "AC-2: the criterion changed after this sheet was made: export the sheet again" in refused
    assert "AC-3: Changes needed needs a Comment, so the author knows why" in refused
    assert "Q-2: Answered by is empty: an answer needs a name" in refused
    assert list(project.read_criteria()) == ["AC-1", "AC-4"]


def test_without_drafts_the_sheet_has_only_cases(make_project, tmp_path):
    root = make_project({"acceptance-criteria.yaml": "criteria:\n  - {id: AC-1, text: t}\n"})
    sheet = root / "review" / "cases.xlsx"
    export_review_sheet(TeamProject.load(root), sheet)
    assert load_workbook(sheet).sheetnames == ["Cases", "How to review"]


def test_an_answer_keeps_the_date_the_sheet_gives_or_today(make_project, tmp_path):
    project, sheet = sheet_with(make_project, tmp_path)
    fill(sheet, "Questions", {2: {"Answer": "Yes", "Answered by": "product-owner", "Date": "2026-10-20"},
                              3: {"Answer": "It retries", "Answered by": "payments-lead"}})
    import_review_sheet(project, sheet)
    first, second = CriteriaDrafts.load(project).questions
    assert first["answered_on"] == "2026-10-20"
    assert second["answered_on"] == datetime.date.today().isoformat()


def test_an_answer_with_a_date_that_isnt_one_is_refused(make_project, tmp_path):
    project, sheet = sheet_with(make_project, tmp_path)
    fill(sheet, "Questions", {2: {"Answer": "Yes", "Answered by": "product-owner", "Date": "20/10/2026"}})
    refused = "\n".join(import_review_sheet(project, sheet).refused)
    assert "Q-1: Date '20/10/2026' (write it as 2026-10-24, or leave it empty for today)" in refused
    assert [q["id"] for q in CriteriaDrafts.load(project).open_questions()] == ["Q-1", "Q-2"]
