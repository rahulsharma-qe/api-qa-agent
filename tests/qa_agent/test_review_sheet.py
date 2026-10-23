"""The review sheet: export, a person's decisions, import with every guard."""

from openpyxl import load_workbook

from qa_agent.cases import load_case_file
from qa_agent.cli import main
from qa_agent.project import TeamProject
from qa_agent.review_sheet import export_review_sheet, import_review_sheet

from tests.qa_agent.conftest import CRITERIA

CASES = """
feature: Bookings
cases:
  - id: TC1
    title: Create
    criteria: [AC-1]
    steps:
      - request: {method: POST, path: /bookings, body: {amount: 250}}
        expect: {status: 201, fields: {status: active}}
        capture: {pnr: pnr}
  - id: TC2
    title: In words
    criteria: [AC-1]
    manual_steps: [Make a booking]
  - id: TC3
    title: Broken
    criteria: [AC-9]
    steps:
      - request: {method: GET, path: /x}
        expect: {status: 200}
  - id: TC4
    title: Done already
    criteria: [AC-1]
    manual_steps: [x]
    review: {status: rejected, comment: no}
"""


def exported(make_project):
    root = make_project({"acceptance-criteria.yaml": CRITERIA, "cases/bookings.yaml": CASES})
    project = TeamProject.load(root)
    sheet = root / "review/cases.xlsx"
    return project, sheet, export_review_sheet(project, sheet)


def decide(sheet, decisions: dict, edits: dict = None):
    book = load_workbook(sheet)
    rows = book["Cases"]
    header = [cell.value for cell in rows[1]]
    for row in range(2, rows.max_row + 1):
        case_id = rows.cell(row, 1).value
        for name, value in {**(edits or {}).get(case_id, {}), **decisions.get(case_id, {})}.items():
            rows.cell(row, header.index(name) + 1).value = value
    book.save(sheet)


def test_export_lists_cases_to_review_with_readable_locked_steps(make_project):
    project, sheet, count = exported(make_project)
    rows = load_workbook(sheet)["Cases"]
    assert count == 3                                                # TC4 was rejected: not in the sheet
    header = [cell.value for cell in rows[1]]
    first = {name: cell for name, cell in zip(header, rows[2])}
    assert first["Steps"].value == '1. POST /bookings {"amount": 250} → remember pnr'
    assert first["Expected"].value == '1. 201; status = "active"'
    assert first["Steps"].protection.locked and not first["Decision"].protection.locked
    assert rows.protection.sheet and rows.column_dimensions["P"].hidden
    second = {name: cell.value for name, cell in zip(header, rows[3])}
    assert (second["ID"], second["Expected"]) == ("TC2", "(in words: needs steps; no expected result written)")
    third = {name: cell.value for name, cell in zip(header, rows[4])}
    assert "criterion AC-9 is not in acceptance-criteria.yaml" in third["Check findings"]


def test_rows_are_in_id_order_and_a_case_in_words_shows_its_expected_result(make_project):
    root = make_project({"cases/b.yaml": """
        feature: B
        cases:
          - {id: TC10, title: ten, manual_steps: [Open it]}
          - {id: TC9, title: nine, manual_steps: [Open it, "Expected: it opens"]}
          - {id: TC2, title: two, manual_steps: [Given x, Then it is refused]}
        """})
    sheet = root / "review/cases.xlsx"
    export_review_sheet(TeamProject.load(root), sheet)
    rows = list(load_workbook(sheet)["Cases"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == ["TC2", "TC9", "TC10"]
    assert (rows[1][8], rows[1][9]) == ("Open it", "it opens\n(in words: needs steps)")
    assert (rows[0][8], rows[0][9]) == ("Given x", "Then it is refused\n(in words: needs steps)")


def test_decisions_are_applied_with_who_and_when_and_edits(make_project):
    project, sheet, _ = exported(make_project)
    decide(sheet, {"TC1": {"Decision": "Approve", "Reviewer": "qa-lead", "Date": "2026-10-21"}},
           edits={"TC1": {"Title": "Create a booking", "Priority": "P0", "Tags": "smoke"}})
    result = import_review_sheet(project, sheet)
    case = load_case_file(project.cases_dir / "bookings.yaml").cases[0]
    assert result.applied == ["TC1: approved by qa-lead, 2026-10-21"] and not result.refused
    assert (case.title, case.priority, case.tags) == ("Create a booking", "P0", ["smoke"])
    assert case.is_approved and case.review.by == "qa-lead" and case.review.date == "2026-10-21"


def test_cases_in_words_or_with_errors_cant_be_approved(make_project):
    project, sheet, _ = exported(make_project)
    decide(sheet, {"TC2": {"Decision": "Approve", "Reviewer": "qa-lead"},
                   "TC3": {"Decision": "Approve", "Reviewer": "qa-lead"}})
    refused = "\n".join(import_review_sheet(project, sheet).refused)
    assert "TC2: can't be approved: it has steps only in words" in refused
    assert "TC3: can't be approved: check finds 1 error(s)" in refused


def test_a_case_changed_after_export_is_not_applied(make_project):
    project, sheet, _ = exported(make_project)
    path = project.cases_dir / "bookings.yaml"
    path.write_text(path.read_text().replace("amount: 250", "amount: 1"))
    decide(sheet, {"TC1": {"Decision": "Approve", "Reviewer": "qa-lead"}})
    result = import_review_sheet(project, sheet)
    assert result.refused == ["TC1: the case changed after this sheet was made: export the sheet again"]
    assert load_case_file(path).cases[0].review.status == "draft"


def test_a_decision_needs_a_reviewer_and_a_send_back_needs_a_comment(make_project):
    project, sheet, _ = exported(make_project)
    decide(sheet, {"TC1": {"Decision": "Approve"}, "TC3": {"Decision": "Changes needed", "Reviewer": "qa-lead"}})
    refused = "\n".join(import_review_sheet(project, sheet).refused)
    assert "TC1: Reviewer is empty" in refused and "TC3: Changes needed needs a Comment" in refused


def test_the_commands_and_their_exit_codes(make_project, capsys):
    root = make_project({"acceptance-criteria.yaml": CRITERIA, "cases/bookings.yaml": CASES})
    assert main(["--project", str(root), "review", "apply"]) == 2                  # no sheet yet
    assert main(["--project", str(root), "review", "export"]) == 0
    decide(root / "review/cases.xlsx", {"TC3": {"Decision": "Reject", "Reviewer": "qa-lead", "Comment": "dup"}})
    assert main(["--project", str(root), "review", "apply"]) == 0
    decide(root / "review/cases.xlsx", {"TC2": {"Decision": "Approve", "Reviewer": "qa-lead"}})
    assert main(["--project", str(root), "review", "apply"]) == 1
    assert "refused" in capsys.readouterr().out


def test_a_sheet_with_a_column_removed_is_a_setup_problem(make_project, capsys):
    project, sheet, _ = exported(make_project)
    book = load_workbook(sheet)
    book["Cases"].delete_cols(13)                                    # Reviewer
    book.save(sheet)
    assert main(["--project", str(project.root), "review", "apply"]) == 2
    assert "columns missing: Reviewer" in capsys.readouterr().err
