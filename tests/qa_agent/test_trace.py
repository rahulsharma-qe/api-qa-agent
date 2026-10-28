"""qa-agent trace: the matrix, the coverage rules, results from JUnit, and --check."""

from qa_agent.cases import load_case_file, save_case_file
from qa_agent.cli import main
from qa_agent.generate import generate
from qa_agent.project import TeamProject
from qa_agent.trace import build_trace

API = """
services:
  bookings:
    operations:
      - {method: POST, path: /bookings, responses: [201, 401, 422]}
      - {method: GET, path: "/bookings/{pnr}", responses: [200, 401, 404, 422]}
      - {method: GET, path: /bookings, responses: [200, 401]}
      - {method: POST, path: /test/reset, responses: [204], not_tested: a test-only route}
"""
CRITERIA = """
criteria:
  - {id: AC-1, text: make}
  - {id: AC-2, text: read}
  - {id: AC-3, text: list}
"""
CASES = """
feature: Bookings
cases:
  - id: TC1
    title: Make one
    criteria: [AC-1]
    steps:
      - request: {method: POST, path: /bookings, body: {amount: 1}}
        expect: {status: 201}
  - id: TC2
    title: Read by status
    criteria: [AC-2]
    data_file: data/reads.csv
    steps:
      - request: {method: GET, path: "/bookings/{{row.pnr}}"}
        expect: {status: "{{row.status}}"}
  - id: TC10
    title: In words
    criteria: [AC-2]
    manual_steps: [List them]
"""
BOTH = "team: t\nservice: bookings\nframeworks: [pytest, restassured]\n"


def project_with_tests(make_project):
    root = make_project({"api/api.yaml": API, "acceptance-criteria.yaml": CRITERIA, "cases/bookings.yaml": CASES,
                         "data/reads.csv": "pnr,status\nAB12CD,200\nNOPE99,404\n"}, config=BOTH)
    case_file = load_case_file(root / "cases/bookings.yaml")
    for case in case_file.cases[:2]:
        case.review.status, case.review.by, case.review.fingerprint = "approved", "qa-lead", case.fingerprint()
    save_case_file(case_file)
    project = TeamProject.load(root)
    generate(project)
    return project


def junit(path, outcomes):
    cases = "".join(f'<testcase classname="c" name="{name}">{body}</testcase>' for name, body in outcomes)
    path.write_text(f'<?xml version="1.0"?><testsuites><testsuite name="s">{cases}</testsuite></testsuites>')
    return path


def test_coverage_follows_the_rules(make_project):
    report = build_trace(project_with_tests(make_project))
    coverage = {item.name: item for item in report.coverage}
    assert (coverage["Criteria"].covered, coverage["Criteria"].total) == (2, 3)
    assert coverage["Criteria"].missing == ["AC-3 (no case)"]
    assert (coverage["Operations"].covered, coverage["Operations"].total) == (2, 3)     # /test/reset left out
    assert coverage["Operations"].missing == ["GET /bookings"]
    codes = coverage["Response codes"]
    assert (codes.covered, codes.total) == (3, 6)                  # 201, and 200 and 404 from the data rows
    assert codes.missing == ["POST /bookings 422", "GET /bookings 200", "401 (bookings, any operation)"]
    assert report.not_tested == ["POST /test/reset: a test-only route"]
    assert (coverage["Cases to tests"].covered, coverage["Cases to tests"].total) == (2, 2)


def test_the_matrix_lists_every_criterion_and_case_with_its_tests(make_project):
    report = build_trace(project_with_tests(make_project))
    rows = [(row.criterion, row.case_id, row.review) for row in report.rows]
    assert rows == [("AC-1", "TC1", "approved"), ("AC-2", "TC2", "approved"), ("AC-2", "TC10", "draft (needs steps)"),
                    ("AC-3", "", "-")]
    assert report.rows[0].tests == ["pytest: test_bookings.py::test_tc1_make_one",
                                    "restassured: BookingsTest.java::tc1MakeOne"]


def test_results_come_from_junit_per_framework_with_a_run_per_data_row(make_project, tmp_path):
    project = project_with_tests(make_project)
    python = junit(tmp_path / "junit.xml", [("test_tc1_make_one", ""), ("test_tc2_read_by_status[0]", ""),
                                            ("test_tc2_read_by_status[1]", "<failure message='x'/>")])
    surefire = tmp_path / "surefire"
    surefire.mkdir()
    junit(surefire / "TEST-TestSuite.xml", [("tc1MakeOne", "<skipped/>")])
    (surefire / "testng-results.xml").write_text("<testng-results/>")
    report = build_trace(project, [python, surefire])
    results = {row.case_id: row.result for row in report.rows if row.case_id}
    assert results["TC1"] == "pytest passed 1/1 · restassured skipped"
    assert results["TC2"] == "pytest failed 1/2 · restassured not run"
    assert report.results_read == {"pytest": {"passed": 2, "failed": 1}, "restassured": {"skipped": 1}}


def test_a_missing_test_and_a_test_without_a_case_are_listed_and_check_fails(make_project, capsys):
    project = project_with_tests(make_project)
    (project.root / "src/test/java/generated/BookingsTest.java").unlink()
    python = project.root / "tests/generated/test_bookings.py"
    python.write_text(python.read_text() + "\n\ndef test_tc9_left_over(api):\n    pass\n")
    report = build_trace(project)
    assert report.missing_tests == ["TC1 (restassured)", "TC2 (restassured)"]
    assert report.orphan_tests == ["pytest: test_bookings.py::test_tc9_left_over"]
    assert main(["--project", str(project.root), "trace"]) == 0
    assert main(["--project", str(project.root), "trace", "--check"]) == 1
    assert (project.root / "reports/trace.md").is_file() and (project.root / "reports/trace.xlsx").is_file()
    assert "TC1 (restassured), TC2 (restassured): run qa-agent generate" in capsys.readouterr().out


def test_results_that_dont_exist_are_a_setup_problem(make_project, capsys):
    project = project_with_tests(make_project)
    assert main(["--project", str(project.root), "trace", "--results", "reports/nothing.xml"]) == 2
    assert "no such results: reports/nothing.xml" in capsys.readouterr().err


def test_a_coverage_floor_fails_only_when_given_and_missed(make_project, capsys):
    project = project_with_tests(make_project)
    assert main(["--project", str(project.root), "trace", "--min-coverage", "60"]) == 0   # 67% and 67%
    assert main(["--project", str(project.root), "trace", "--min-coverage", "70"]) == 1
    assert "Criteria 67% (floor 70%), Operations 67% (floor 70%)" in capsys.readouterr().out


def test_a_rejected_case_is_not_listed_as_waiting_for_approval(make_project):
    from qa_agent.trace import _criteria_coverage
    from qa_agent.cases import Case, Review
    from qa_agent.project import Criterion
    criteria = {"AC-1": Criterion("AC-1", "a"), "AC-2": Criterion("AC-2", "b")}
    cases = [Case("TC1", "draft one", criteria=["AC-1"]),
             Case("TC2", "rejected one", criteria=["AC-1"], review=Review(status="rejected")),
             Case("TC3", "rejected only", criteria=["AC-2"], review=Review(status="rejected"))]
    assert _criteria_coverage(criteria, cases).missing == ["AC-1 (not approved: TC1)", "AC-2 (no case)"]


def test_results_that_cant_be_read_or_hold_no_tests_stop_the_trace(make_project, tmp_path, capsys):
    project = project_with_tests(make_project)
    broken = tmp_path / "broken.xml"
    broken.write_text("<testsuite><testcase name=")
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main(["--project", str(project.root), "trace", "--check", "--results", str(broken)]) == 2
    assert "results can't be read (not valid XML)" in capsys.readouterr().err
    assert main(["--project", str(project.root), "trace", "--check", "--results", str(empty)]) == 2
    assert "no JUnit test results in" in capsys.readouterr().err


def test_results_of_tests_that_no_longer_exist_are_left_out(make_project, tmp_path):
    project = project_with_tests(make_project)
    old = junit(tmp_path / "junit.xml", [("test_tc1_make_one", ""), ("test_tc9_removed_since", "")])
    report = build_trace(project, [old])
    assert report.results_read == {"pytest": {"passed": 1}} and report.results_left_out == 1
    assert "1 result of tests that no longer exist left out" in report.as_text()
    markdown = report.as_markdown()
    assert "Results: pytest 1 passed · 1 result of tests that no longer exist left out" in markdown
    assert "Covered means an approved case exists" in markdown
    report.write_excel(tmp_path / "trace.xlsx")
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "trace.xlsx")["Coverage"]
    shown = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("no longer exist left out" in value for value in shown)
    assert any(value.startswith("Covered means") for value in shown)


def test_results_only_of_tests_that_no_longer_exist_stop_the_trace(make_project, tmp_path, capsys):
    project = project_with_tests(make_project)
    old = junit(tmp_path / "junit.xml", [("test_tc8_gone", ""), ("test_tc9_removed_since", "")])
    assert main(["--project", str(project.root), "trace", "--check", "--results", str(old)]) == 2
    assert "none of the results are of tests that exist now (2 results" in capsys.readouterr().err
